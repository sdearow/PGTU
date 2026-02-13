#!/usr/bin/env python3
"""Convert two Excel files into JavaScript data files."""

import json
import openpyxl

def clean_value(val):
    """Convert cell value to a clean Python value."""
    if val is None:
        return None
    if isinstance(val, float):
        if val == int(val):
            return int(val)
        return round(val, 2)
    if isinstance(val, int):
        return val
    if isinstance(val, str):
        stripped = val.strip()
        if stripped == "":
            return None
        return stripped
    return str(val)

def sheet_to_records(ws, headers):
    """Convert a worksheet to a list of dicts using provided headers."""
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_vals = list(row)[:len(headers)]
        cleaned = [clean_value(v) for v in row_vals]
        if all(v is None for v in cleaned):
            continue
        record = {}
        for h, v in zip(headers, cleaned):
            record[h] = v if v is not None else ""
        records.append(record)
    return records

def convert_file(filepath, expected_headers, var_name, output_path):
    """Convert an Excel workbook into a JS data file."""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    data = {}

    for i in range(1, 16):
        sheet_name = f"Municipio {i}"
        if sheet_name not in wb.sheetnames:
            print(f"  WARNING: Sheet '{sheet_name}' not found in {filepath}")
            data[str(i)] = []
            continue

        ws = wb[sheet_name]
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        actual_headers = [str(h).strip() if h is not None else "" for h in first_row]
        while actual_headers and actual_headers[-1] == "":
            actual_headers.pop()

        print(f"  Sheet '{sheet_name}': {len(actual_headers)} columns -> {actual_headers}")

        headers = actual_headers
        records = sheet_to_records(ws, headers)
        data[str(i)] = records
        print(f"    -> {len(records)} records")

    wb.close()

    json_str = json.dumps(data, ensure_ascii=False, indent=2)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(f"var {var_name} = {json_str};\n")

    print(f"  Written: {output_path}")

# --- Tabella 1 ---
print("=== Converting Tabella 1 ===")
headers_t1 = [
    "Identificazione Strada", "Is_PGTU", "Classificazione", "% PGTU",
    "Flusso medio", "Flusso Min", "Flusso Max", "Is_TPL", "Is_GV", "Municipio"
]
convert_file(
    "/home/user/PGTU/Mappe Proposte_130224/Tabella 1/Tabella_Sintesi_1_v2.xlsx",
    headers_t1,
    "dataTabella1",
    "/home/user/PGTU/data_tabella1.js"
)

# --- Tabella 2 ---
print("\n=== Converting Tabella 2 ===")
headers_t2 = [
    "Identificazione Strada", "Is_PGTU", "Flusso medio", "Flusso Min",
    "Flusso Max", "Is_TPL", "Is_GV", "Municipio"
]
convert_file(
    "/home/user/PGTU/Mappe Proposte_130224/Tabella 2/Tabella_Sintesi_2_050126.xlsx",
    headers_t2,
    "dataTabella2",
    "/home/user/PGTU/data_tabella2.js"
)

print("\nDone!")
